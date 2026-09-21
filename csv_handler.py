"""CSV imports use the same service as the manual form."""
import csv
from datetime import date, datetime
from pathlib import Path

def import_file(path, service):
    suffix = Path(path).suffix.lower()
    if suffix == '.csv':
        return import_csv(path, service)
    if suffix != '.xlsx':
        raise ValueError('Choose a CSV or Excel .xlsx file. Save older .xls files as .xlsx first.')
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ValueError('Excel import requires openpyxl. Run: python -m pip install -r requirement.txt') from None
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
    except Exception as error:
        raise ValueError(f'Cannot read Excel workbook: {error}') from error
    try:
        sheet = workbook.worksheets[0]
        cells = sheet.iter_rows()
        headers = [str(c.value).strip() if c.value is not None else '' for c in next(cells, [])]
        while headers and not headers[-1]:
            headers.pop()
        if not headers or any(not h for h in headers) or len(set(headers)) != len(headers):
            raise ValueError('The first worksheet needs unique, non-empty column names in its first row.')
        def records():
            for row in cells:
                if all(c.value is None for c in row):
                    yield None
                    continue
                if any(c.data_type == 'f' for c in row):
                    yield {'__error__':'Formula cells are not supported. Paste values before importing.'}
                    continue
                if any(c.value is not None for c in row[len(headers):]):
                    yield {'__error__':'Row contains data outside the named columns.'}
                    continue
                values=[]
                for cell in row[:len(headers)]:
                    value=cell.value
                    if isinstance(value, datetime): value=value.date().isoformat()
                    elif isinstance(value, date): value=value.isoformat()
                    elif isinstance(value, bool): value=str(value).lower()
                    values.append('' if value is None else str(value))
                yield dict(zip(headers,values))
        return _import_records(headers, records(), service)
    finally:
        workbook.close()

def import_csv(path, service):
    with open(path, newline='', encoding='utf-8-sig') as source:
        reader = csv.DictReader(source)
        return _import_records(reader.fieldnames, reader, service)

def _import_records(headers, records, service):
        accepted, issues = 0, []
        needed = {'invoice_id','supplier_id','supplier_name','invoice_number','invoice_date','invoice_due_date','invoice_amount','currency'}
        if not needed.issubset(headers or []):
            raise ValueError('File is missing required columns: '+', '.join(sorted(needed-set(headers or []))))
        for line, row in enumerate(records,2):
            if row is None:
                continue
            try:
                if '__error__' in row:
                    raise ValueError(row['__error__'])
                if None in row or any(v is None for v in row.values()):
                    raise ValueError('CSV row has a different number of fields than its header.')
                service.add_invoice(row)
                accepted += 1
                if row.get('payment_date') and row.get('days_to_payment') and row.get('invoice_submission_date'):
                    expected = (date.fromisoformat(row['payment_date'])-date.fromisoformat(row['invoice_submission_date'])).days
                    if str(expected)!=row['days_to_payment']:
                        issues.append({'line':line,'invoice_id':row['invoice_id'],'result':'Imported with warning','reason':f'Source days_to_payment is {row["days_to_payment"]}; calculated value is {expected}.'})
            except ValueError as error:
                issues.append({'line':line,'invoice_id':row.get('invoice_id',''),'result':'Rejected','reason':str(error)})
        return accepted, issues

def export_csv(path, rows):
    if not rows:
        raise ValueError('There are no rows to export.')
    with open(path,'w',newline='',encoding='utf-8-sig') as target:
        writer = csv.DictWriter(target,fieldnames=list(rows[0]))
        writer.writeheader()
        for row in rows:
            # Prevent spreadsheet formula execution when exported text is opened in Excel.
            writer.writerow({k:("'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v) for k,v in row.items()})
